import math
import os
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
READY_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
REJECT_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
BORDER = Border(*(Side(style="thin", color="D9D9D9"),) * 4)
MONEY_COLS = {"Billed_Amount", "Allowed_Amount", "Copay", "Deductible", "Patient_Responsibility", "Insurance_Payable"}


def _safe(v):
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return None
    if isinstance(v, str):
        v = ILLEGAL_CHARACTERS_RE.sub("", v)[:32000]
    return v


def _write_table(ws, df):
    ws.append(list(df.columns))
    for c in ws[1]:
        c.fill, c.font, c.border = HEADER_FILL, Font(bold=True, color="FFFFFF"), BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    status_idx = list(df.columns).index("Validation_Status") + 1 if "Validation_Status" in df.columns else None
    for row in df.itertuples(index=False, name=None):
        ws.append([_safe(v) for v in row])
    for r in ws.iter_rows(min_row=2):
        fill = None
        if status_idx:
            fill = READY_FILL if r[status_idx - 1].value == config.STATUS_READY else REJECT_FILL
        for c in r:
            c.border = BORDER
            if isinstance(c.value, str) and c.value.startswith("="):
                c.data_type = "s"   # text ko formula na banne do
            if fill:
                c.fill = fill
            if ws.cell(1, c.column).value in MONEY_COLS and isinstance(c.value, (int, float)):
                c.number_format = '#,##0.00'
    for i, col in enumerate(df.columns, 1):
        longest = max([len(str(col))] + [len(str(v)) for v in df[col].astype(str).head(500)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 12), 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _summary_rows(df):
    total = len(df)
    ready = int((df["Validation_Status"] == config.STATUS_READY).sum())
    rows = [
        ("Report generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Total claims", total),
        ("Ready to submit", ready),
        ("Rejected", total - ready),
        ("Clean claim rate %", round(ready / total * 100, 1) if total else 0),
        ("Total billed (all)", round(float(df["Billed_Amount"].sum()), 2)),
        ("Total billed (ready)", round(float(df.loc[df["Validation_Status"] == config.STATUS_READY, "Billed_Amount"].sum()), 2)),
        ("Expected insurance payable", round(float(df["Insurance_Payable"].sum()), 2)),
        ("Expected patient responsibility", round(float(df["Patient_Responsibility"].sum()), 2)),
        ("", ""),
        ("TOP REJECTION REASONS", "Count"),
    ]
    reasons = {}
    for text in df.loc[df["Validation_Status"] == config.STATUS_REJECTED, "Validation_Issues"]:
        for part in str(text).split("; "):
            reason = part.split(" (same as row")[0]
            reasons[reason] = reasons.get(reason, 0) + 1
    rows += sorted(reasons.items(), key=lambda kv: -kv[1])[:10] or [("None", 0)]
    return rows


def export_audit_report(df, output_path):
    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Summary"
    for label, value in _summary_rows(df):
        ws_sum.append([label, value])
    for c in ws_sum["A"]:
        c.font = Font(bold=True)
    ws_sum.column_dimensions["A"].width = 45
    ws_sum.column_dimensions["B"].width = 22

    _write_table(wb.create_sheet("Audit_Report"), df)
    rejected = df[df["Validation_Status"] == config.STATUS_REJECTED]
    if len(rejected):
        _write_table(wb.create_sheet("Rejected_Claims"), rejected)

    try:
        wb.save(output_path)
    except PermissionError:
        base, ext = os.path.splitext(output_path)
        output_path = f"{base}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{ext}"
        wb.save(output_path)
        print("NOTE: Purani report Excel mein khuli thi, naye naam se save ki.")
    return output_path
